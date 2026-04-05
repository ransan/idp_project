import os
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

from app.schemas import ParseResult

# Maximum characters to keep when truncating long documents
MAX_TEXT_LENGTH = 50_000
# Minimum characters from text extraction before falling back to OCR
OCR_FALLBACK_THRESHOLD = 50


class ParserError(Exception):
    pass


class UnsupportedFileTypeError(ParserError):
    pass


class CorruptedFileError(ParserError):
    pass


class PasswordProtectedError(ParserError):
    pass


class DocumentParser:
    MIME_MAP = {
        "application/pdf": "_parse_pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "_parse_docx",
        "image/png": "_parse_image",
        "image/jpeg": "_parse_image",
        "image/jpg": "_parse_image",
        "image/tiff": "_parse_image",
    }

    EXT_MIME_MAP = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".tiff": "image/tiff",
    }

    def parse(self, file_path: str, mime_type: str | None = None) -> ParseResult:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if not mime_type:
            mime_type = self.EXT_MIME_MAP.get(path.suffix.lower())

        if not mime_type or mime_type not in self.MIME_MAP:
            raise UnsupportedFileTypeError(
                f"Unsupported file type: {mime_type or path.suffix}"
            )

        handler = getattr(self, self.MIME_MAP[mime_type])
        try:
            result = handler(file_path)
        except (PasswordProtectedError, UnsupportedFileTypeError):
            raise
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse file: {e}") from e

        # Truncate very long documents preserving beginning and end
        if len(result.text) > MAX_TEXT_LENGTH:
            half = MAX_TEXT_LENGTH // 2
            result = ParseResult(
                text=result.text[:half] + "\n\n[... content truncated ...]\n\n" + result.text[-half:],
                page_count=result.page_count,
                method=result.method + "+truncated",
            )

        return result

    def _parse_pdf(self, file_path: str) -> ParseResult:
        try:
            doc = fitz.open(file_path)
        except fitz.fitz.FileDataError as e:
            raise CorruptedFileError(f"Corrupted PDF: {e}") from e

        if doc.is_encrypted:
            doc.close()
            raise PasswordProtectedError("PDF is password-protected")

        page_count = len(doc)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()

        text = "\n".join(text_parts).strip()

        # Fallback to OCR if text extraction yields too little
        if self._needs_ocr(text):
            return self._ocr_pdf(file_path, page_count)

        return ParseResult(text=text, page_count=page_count, method="pymupdf")

    def _ocr_pdf(self, file_path: str, page_count: int) -> ParseResult:
        doc = fitz.open(file_path)
        text_parts = []
        for page in doc:
            pix = page.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            page_text = pytesseract.image_to_string(img)
            text_parts.append(page_text)
        doc.close()
        text = "\n".join(text_parts).strip()
        return ParseResult(text=text, page_count=page_count, method="pymupdf+ocr")

    def _parse_docx(self, file_path: str) -> ParseResult:
        doc = DocxDocument(file_path)
        text_parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(cells))
            if table_rows:
                text_parts.append("\n".join(table_rows))

        text = "\n".join(text_parts).strip()
        # DOCX doesn't have a native page count; estimate from content
        estimated_pages = max(1, len(text) // 3000)
        return ParseResult(text=text, page_count=estimated_pages, method="python-docx")

    def _parse_image(self, file_path: str) -> ParseResult:
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img).strip()
        return ParseResult(text=text, page_count=1, method="tesseract")

    @staticmethod
    def _needs_ocr(text: str) -> bool:
        return len(text.strip()) < OCR_FALLBACK_THRESHOLD
