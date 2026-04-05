import asyncio
import io
import os
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import String, event
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.types import JSON, TypeDecorator

from app.auth import generate_api_key, hash_api_key, hash_password, create_access_token
from app.models import APIKey, Base, Client, DocStatus, Document, User

# ---------------------------------------------------------------------------
# SQLite compatibility: Replace PostgreSQL-specific types for testing
# ---------------------------------------------------------------------------
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.dialects.postgresql import UUID as PG_UUID


class SQLiteUUID(TypeDecorator):
    """UUID type that stores as String in SQLite but handles UUID objects transparently."""
    impl = String
    cache_ok = True

    def process_bind_param(self, value, dialect):
        if value is not None:
            return str(value)
        return value

    def process_result_value(self, value, dialect):
        if value is not None:
            return uuid.UUID(value) if not isinstance(value, uuid.UUID) else value
        return value


def _make_sqlite_compatible():
    """Replace JSONB -> JSON and UUID -> SQLiteUUID for SQLite testing."""
    from sqlalchemy import Enum as SAEnum

    for table in Base.metadata.tables.values():
        for column in table.columns:
            if isinstance(column.type, JSONB):
                column.type = JSON()
            if isinstance(column.type, PG_UUID):
                column.type = SQLiteUUID(36)
            # SQLite needs native_enum=False for enums
            if isinstance(column.type, SAEnum):
                column.type = SAEnum(
                    *column.type.enums,
                    name=column.type.name,
                    native_enum=False,
                )


_make_sqlite_compatible()


# ---------------------------------------------------------------------------
# Async event loop
# ---------------------------------------------------------------------------
@pytest.fixture(scope="session")
def event_loop():
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


# ---------------------------------------------------------------------------
# Test database (async SQLite in-memory)
# ---------------------------------------------------------------------------
TEST_DB_URL = "sqlite+aiosqlite:///:memory:"


@pytest_asyncio.fixture
async def db_engine():
    engine = create_async_engine(TEST_DB_URL, echo=False)

    # Register UUID-to-string conversion for SQLite
    @event.listens_for(engine.sync_engine, "connect")
    def _set_sqlite_pragma(dbapi_conn, connection_record):
        import sqlite3

        # Register UUID adapter
        sqlite3.register_adapter(uuid.UUID, lambda u: str(u))

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield engine
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
    await engine.dispose()


@pytest_asyncio.fixture
async def db_session(db_engine):
    session_factory = async_sessionmaker(db_engine, class_=AsyncSession, expire_on_commit=False)
    async with session_factory() as session:
        yield session


# ---------------------------------------------------------------------------
# Auth seed data
# ---------------------------------------------------------------------------
TEST_CLIENT_ID = "test-client"
TEST_RAW_API_KEY = "dp_live_000000000000000000000000deadbeef"
TEST_ADMIN_RAW_KEY = "dp_live_000000000000000000000000cafebabe"
TEST_USER_EMAIL = "testuser@example.com"
TEST_USER_PASSWORD = "testpassword123"


async def _seed_auth_data(session: AsyncSession):
    """Create a Client, API keys (member + admin), and a User for tests."""
    client_record = Client(
        id=TEST_CLIENT_ID,
        name="Test Client",
        email="client@example.com",
        plan="free",
        rate_limit=1000,
    )
    session.add(client_record)

    member_key = APIKey(
        key_hash=hash_api_key(TEST_RAW_API_KEY),
        key_prefix=TEST_RAW_API_KEY[:12],
        client_id=TEST_CLIENT_ID,
        name="Test Key",
        role="member",
    )
    session.add(member_key)

    admin_key = APIKey(
        key_hash=hash_api_key(TEST_ADMIN_RAW_KEY),
        key_prefix=TEST_ADMIN_RAW_KEY[:12],
        client_id=TEST_CLIENT_ID,
        name="Admin Key",
        role="admin",
    )
    session.add(admin_key)

    user = User(
        client_id=TEST_CLIENT_ID,
        email=TEST_USER_EMAIL,
        password_hash=hash_password(TEST_USER_PASSWORD),
        name="Test User",
        role="admin",
    )
    session.add(user)

    await session.commit()


# ---------------------------------------------------------------------------
# FastAPI test client
# ---------------------------------------------------------------------------
@pytest_asyncio.fixture
async def client(db_engine):
    """Create a test client with overridden DB dependency, mocked Celery tasks, and auth seed data."""
    session_factory = async_sessionmaker(db_engine, class_=AsyncSession, expire_on_commit=False)

    # Seed auth data
    async with session_factory() as seed_session:
        await _seed_auth_data(seed_session)

    async def override_get_db():
        async with session_factory() as session:
            try:
                yield session
                await session.commit()
            except Exception:
                await session.rollback()
                raise

    # Patch Celery tasks before importing app
    with (
        patch("app.routers.documents.process_document_task") as mock_doc_task,
        patch("app.routers.documents.process_batch_task") as mock_batch_task,
    ):
        mock_doc_task.delay = MagicMock(return_value=MagicMock(id="test-task-id"))
        mock_batch_task.delay = MagicMock(return_value=MagicMock(id="test-batch-id"))

        from app.database import get_db
        from app.main import app

        app.dependency_overrides[get_db] = override_get_db

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as ac:
            yield ac

        app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# Auth header helpers
# ---------------------------------------------------------------------------
@pytest.fixture
def auth_headers():
    """Headers with a member API key."""
    return {"X-API-Key": TEST_RAW_API_KEY}


@pytest.fixture
def admin_headers():
    """Headers with an admin API key."""
    return {"X-API-Key": TEST_ADMIN_RAW_KEY}


# ---------------------------------------------------------------------------
# Temp directory for uploads
# ---------------------------------------------------------------------------
@pytest.fixture
def tmp_upload_dir(tmp_path):
    with patch("app.config.settings.UPLOAD_DIR", str(tmp_path)):
        yield tmp_path


# ---------------------------------------------------------------------------
# Sample file fixtures
# ---------------------------------------------------------------------------
@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    """Create a minimal valid PDF file."""
    import fitz

    pdf_path = tmp_path / "test.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Invoice #12345\nVendor: Acme Corp\nTotal: $1,500.00\nDue: 2026-05-01")
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_docx(tmp_path) -> Path:
    """Create a minimal valid DOCX file."""
    from docx import Document as DocxDoc

    docx_path = tmp_path / "test.docx"
    doc = DocxDoc()
    doc.add_paragraph("Contract Agreement")
    doc.add_paragraph("Between Party A and Party B")
    doc.add_paragraph("Effective Date: 2026-01-01")

    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Duration"
    table.rows[1].cells[1].text = "12 months"

    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def sample_image(tmp_path) -> Path:
    """Create a minimal test image."""
    from PIL import Image

    img_path = tmp_path / "test.png"
    img = Image.new("RGB", (200, 100), color="white")
    img.save(str(img_path))
    return img_path


@pytest.fixture
def sample_pdf_bytes(sample_pdf) -> bytes:
    return sample_pdf.read_bytes()


@pytest.fixture
def sample_docx_bytes(sample_docx) -> bytes:
    return sample_docx.read_bytes()


@pytest.fixture
def empty_pdf(tmp_path) -> Path:
    """Create an empty PDF (no text)."""
    import fitz

    pdf_path = tmp_path / "empty.pdf"
    doc = fitz.open()
    doc.new_page()  # blank page
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def large_text() -> str:
    """Generate text exceeding MAX_TEXT_LENGTH."""
    return "This is a repeated sentence for testing. " * 5000


@pytest.fixture
def mock_llm_response() -> dict:
    """Standard successful LLM response."""
    return {
        "category": "invoice",
        "confidence": 0.95,
        "summary": "Invoice #12345 from Acme Corp for $1,500 in consulting services.",
        "extracted_data": {
            "vendor": "Acme Corp",
            "buyer": "Test Client",
            "invoice_number": "12345",
            "date": "2026-03-15",
            "due_date": "2026-05-01",
            "total_amount": 1500.00,
            "currency": "USD",
            "payment_terms": "Net 30",
            "line_items": [
                {"description": "Consulting", "quantity": 10, "unit_price": 150, "total": 1500}
            ],
        },
        "key_entities": {
            "people": [],
            "organizations": ["Acme Corp", "Test Client"],
            "dates": ["2026-03-15", "2026-05-01"],
            "monetary_amounts": ["$1,500"],
            "locations": [],
            "reference_numbers": ["12345"],
        },
        "flags": ["Due date is within 30 days"],
    }


@pytest.fixture
def sample_document_in_db():
    """Create a Document instance for unit tests."""

    def _create(**kwargs):
        defaults = {
            "id": uuid.uuid4(),
            "filename": "test_file.pdf",
            "original_name": "invoice.pdf",
            "file_size": 1024,
            "mime_type": "application/pdf",
            "storage_path": "/tmp/test_file.pdf",
            "status": DocStatus.UPLOADED,
            "client_id": "test_client",
            "created_at": datetime.now(timezone.utc),
        }
        defaults.update(kwargs)
        return Document(**defaults)

    return _create
